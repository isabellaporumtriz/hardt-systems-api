from __future__ import annotations

import re
import unicodedata

from io import BytesIO
from uuid import UUID, uuid4

from openpyxl import load_workbook
from sqlalchemy import insert, select
from sqlalchemy.orm import Session

from app.argos.intake import (
    ArgosIntakeError,
    is_valid_cnpj,
)
from app.argos.models import (
    ArgosCnpjBatch,
    ArgosCnpjBatchItem,
)
from app.downloads.storage import (
    delete_private_object,
    upload_private_object,
)


MAX_CNPJ_XLSX_BYTES = 64 * 1024 * 1024


def _cell_text(
    value: object,
) -> str:
    if value is None:
        return ""

    if isinstance(value, float):
        if value.is_integer():
            return str(int(value))

    return str(value).strip()


def _normalize_header(
    value: object,
) -> str:
    raw = _cell_text(
        value
    )

    normalized = unicodedata.normalize(
        "NFKD",
        raw,
    )

    without_accents = "".join(
        char
        for char in normalized
        if not unicodedata.combining(
            char
        )
    )

    return re.sub(
        r"[^a-z0-9]+",
        "_",
        without_accents.lower(),
    ).strip("_")


def _cnpj_digits(
    value: str,
) -> str:
    return re.sub(
        r"\D",
        "",
        value or "",
    )


def _parse_cnpj_workbook(
    content: bytes,
) -> list[dict]:
    try:
        workbook = load_workbook(
            filename=BytesIO(content),
            read_only=True,
            data_only=True,
        )
    except Exception as exc:
        raise ArgosIntakeError(
            "CNPJ_XLSX_READ_FAILED"
        ) from exc

    try:
        worksheet = workbook.active

        if worksheet is None:
            raise ArgosIntakeError(
                "CNPJ_XLSX_NO_SHEET"
            )

        header_cells = next(
            worksheet.iter_rows(
                min_row=1,
                max_row=1,
            ),
            None,
        )

        if not header_cells:
            raise ArgosIntakeError(
                "CNPJ_XLSX_EMPTY"
            )

        headers = [
            _normalize_header(
                cell.value
            )
            for cell in header_cells
        ]

        cnpj_column_index = (
            headers.index("cnpj")
            if "cnpj" in headers
            else None
        )

        receita_headers = {
            "cnpj_basico",
            "cnpj_ordem",
            "cnpj_dv",
        }

        has_receita_columns = (
            receita_headers
            .issubset(set(headers))
        )

        if (
            cnpj_column_index is None
            and not has_receita_columns
        ):
            raise ArgosIntakeError(
                "CNPJ_XLSX_COLUMN_NOT_FOUND"
            )

        if has_receita_columns:
            base_index = headers.index(
                "cnpj_basico"
            )
            ordem_index = headers.index(
                "cnpj_ordem"
            )
            dv_index = headers.index(
                "cnpj_dv"
            )
        else:
            base_index = None
            ordem_index = None
            dv_index = None

        items: list[dict] = []
        seen: set[str] = set()

        for row_number, row in enumerate(
            worksheet.iter_rows(
                min_row=2,
            ),
            start=2,
        ):
            if cnpj_column_index is not None:
                if (
                    cnpj_column_index
                    >= len(row)
                ):
                    continue

                original = _cell_text(
                    row[
                        cnpj_column_index
                    ].value
                )

            else:
                assert base_index is not None
                assert ordem_index is not None
                assert dv_index is not None

                if max(
                    base_index,
                    ordem_index,
                    dv_index,
                ) >= len(row):
                    continue

                base_raw = _cell_text(
                    row[base_index].value
                )
                ordem_raw = _cell_text(
                    row[ordem_index].value
                )
                dv_raw = _cell_text(
                    row[dv_index].value
                )

                if not (
                    base_raw
                    or ordem_raw
                    or dv_raw
                ):
                    continue

                base = _cnpj_digits(
                    base_raw
                )
                ordem = _cnpj_digits(
                    ordem_raw
                )
                dv = _cnpj_digits(
                    dv_raw
                )

                if (
                    base
                    and ordem
                    and dv
                    and len(base) <= 8
                    and len(ordem) <= 4
                    and len(dv) <= 2
                ):
                    combined = (
                        base.zfill(8)
                        + ordem.zfill(4)
                        + dv.zfill(2)
                    )

                    original = (
                        f"{combined[:2]}."
                        f"{combined[2:5]}."
                        f"{combined[5:8]}/"
                        f"{combined[8:12]}-"
                        f"{combined[12:14]}"
                    )

                else:
                    original = (
                        f"{base_raw}/"
                        f"{ordem_raw}-"
                        f"{dv_raw}"
                    )

            # Linhas realmente vazias
            # não pertencem ao lote.
            if not original:
                continue

            digits = _cnpj_digits(
                original
            )

            normalized = (
                digits
                if len(digits) == 14
                else None
            )

            if not is_valid_cnpj(
                original
            ):
                items.append(
                    {
                        "row_number":
                            row_number,
                        "cnpj_original":
                            original,
                        "cnpj_normalized":
                            normalized,
                        "status":
                            "invalid",
                        "error":
                            "CNPJ_INVALID",
                    }
                )
                continue

            assert normalized is not None

            if normalized in seen:
                items.append(
                    {
                        "row_number":
                            row_number,
                        "cnpj_original":
                            original,
                        "cnpj_normalized":
                            normalized,
                        "status":
                            "duplicate",
                        "error":
                            "CNPJ_DUPLICATE_IN_BATCH",
                    }
                )
                continue

            seen.add(
                normalized
            )

            items.append(
                {
                    "row_number":
                        row_number,
                    "cnpj_original":
                        original,
                    "cnpj_normalized":
                        normalized,
                    "status":
                        "pending",
                    "error":
                        None,
                }
            )

        if not items:
            raise ArgosIntakeError(
                "CNPJ_XLSX_NO_CNPJS"
            )

        return items

    finally:
        workbook.close()


def create_cnpj_batch(
    db: Session,
    *,
    original_filename: str,
    mime_type: str,
    content: bytes,
    created_by_user_id: (
        UUID | None
    ),
) -> ArgosCnpjBatch:
    filename = str(
        original_filename
        or "cnpjs.xlsx"
    ).strip()

    if not filename.lower().endswith(
        ".xlsx"
    ):
        raise ArgosIntakeError(
            "CNPJ_FILE_MUST_BE_XLSX"
        )

    if not content:
        raise ArgosIntakeError(
            "CNPJ_FILE_EMPTY"
        )

    if (
        len(content)
        > MAX_CNPJ_XLSX_BYTES
    ):
        raise ArgosIntakeError(
            "CNPJ_FILE_TOO_LARGE"
        )

    parsed_items = (
        _parse_cnpj_workbook(
            content
        )
    )

    valid_count = sum(
        1
        for item in parsed_items
        if item["status"] == "pending"
    )

    duplicate_count = sum(
        1
        for item in parsed_items
        if item["status"]
        == "duplicate"
    )

    invalid_count = sum(
        1
        for item in parsed_items
        if item["status"] == "invalid"
    )

    batch_id = uuid4()

    object_key = (
        "argos/cnpj-batches/"
        f"{batch_id}/"
        "cnpjs.xlsx"
    )

    content_type = (
        "application/"
        "vnd.openxmlformats-officedocument."
        "spreadsheetml.sheet"
    )

    upload_private_object(
        object_key,
        content=content,
        content_type=content_type,
    )

    batch = ArgosCnpjBatch(
        id=batch_id,
        status="parsed",
        original_filename=(
            filename[:255]
        ),
        object_key=object_key,
        mime_type=content_type,
        size_bytes=len(content),
        total_rows=len(
            parsed_items
        ),
        valid_count=valid_count,
        duplicate_count=(
            duplicate_count
        ),
        invalid_count=invalid_count,
        created_by_user_id=(
            created_by_user_id
        ),
    )

    try:
        db.add(
            batch
        )

        # Garante que o batch exista antes
        # dos itens dependentes.
        db.flush()

        chunk_size = 5000

        for start in range(
            0,
            len(parsed_items),
            chunk_size,
        ):
            chunk = parsed_items[
                start:
                start + chunk_size
            ]

            rows = [
                {
                    "id": uuid4(),
                    "batch_id": batch_id,
                    "row_number": (
                        item["row_number"]
                    ),
                    "cnpj_original": (
                        item[
                            "cnpj_original"
                        ][:80]
                    ),
                    "cnpj_normalized": (
                        item[
                            "cnpj_normalized"
                        ]
                    ),
                    "status": (
                        item["status"]
                    ),
                    "error": (
                        item["error"]
                    ),
                    "operation_id": None,
                }
                for item in chunk
            ]

            db.execute(
                insert(
                    ArgosCnpjBatchItem
                ),
                rows,
            )

        db.commit()

        db.refresh(
            batch
        )

    except Exception:
        db.rollback()

        try:
            delete_private_object(
                object_key
            )
        except Exception:
            pass

        raise

    return batch


def get_cnpj_batch_items(
    db: Session,
    *,
    batch_id: UUID,
) -> list[ArgosCnpjBatchItem]:
    return list(
        db.scalars(
            select(
                ArgosCnpjBatchItem
            )
            .where(
                ArgosCnpjBatchItem.batch_id
                == batch_id
            )
            .order_by(
                ArgosCnpjBatchItem.row_number
            )
        ).all()
    )
